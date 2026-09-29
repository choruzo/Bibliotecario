import hashlib
import io
import threading
import time
import uuid
import zipfile

import pytest
from docx import Document as WordDocument
from sqlalchemy import select

from test_h1 import application, login, settings  # noqa: F401
from bibliotecario.converters import convert, edited_provenance, provenance
from bibliotecario.jobs import claim, finalize_conversion, finish_cancel, process, process_delete, renew
from bibliotecario.models import Document, DocumentFile, DocumentVersion, IngestionJob, NormalizedRevision
from bibliotecario.storage import InvalidFile, save_upload, storage_path


def headers(client):
    return {"Origin": "http://localhost:3000", "X-CSRF-Token": client.get("/auth/me").json()["csrf_token"]}


def upload(client, name="document.md", content=b"# Titulo\n\nTexto original.\n", key=None, path="/admin/documents", mime="text/markdown"):
    return client.post(path, files={"file": (name, content, mime)},
                       headers=headers(client) | {"Idempotency-Key": key or str(uuid.uuid4())})


def complete(app, sessions, response):
    lease = claim(sessions, app.state.settings, "test-worker")
    assert lease
    with sessions() as db:
        file = db.scalar(select(DocumentFile).where(DocumentFile.version_id == response["document"]["versions"][0]["id"]))
        result = convert(storage_path(app.state.settings, file.storage_key), file.format, app.state.settings)
    assert finalize_conversion(sessions, app.state.settings, lease, result)
    return result


def test_upload_idempotency_permissions_and_immutable_original(application):
    app, client, sessions = application
    login(client, "lector")
    assert upload(client).status_code == 403
    login(client)
    key = str(uuid.uuid4())
    response = upload(client, key=key)
    assert response.status_code == 201
    data = response.json()
    repeated = upload(client, key=key)
    assert repeated.json()["document"]["id"] == data["document"]["id"]
    assert upload(client, content=b"Different", key=key).status_code == 409
    complete(app, sessions, data)
    version = data["document"]["versions"][0]
    original = client.get(f"/admin/versions/{version['id']}/original")
    assert original.status_code == 200 and original.content == b"# Titulo\n\nTexto original.\n"
    assert "attachment" in original.headers["content-disposition"]
    with sessions() as db:
        assert len(db.scalars(select(Document)).all()) == 1
        assert len(db.scalars(select(DocumentFile)).all()) == 1
        file = db.scalar(select(DocumentFile))
        assert hashlib.sha256(storage_path(app.state.settings, file.storage_key).read_bytes()).hexdigest() == file.sha256


def test_review_history_conflicts_and_manual_provenance(application):
    app, client, sessions = application
    login(client)
    data = upload(client).json()
    complete(app, sessions, data)
    vid = data["document"]["versions"][0]["id"]
    original = client.get(f"/admin/versions/{vid}/normalized").json()
    metadata = data["document"]["versions"][0]["metadata"]
    metadata = metadata | {"author": "Autor revisado", "tags": ["uno", "dos"]}
    edited = "# Titulo\n\nTexto revisado.\n"
    body = {"expected_revision_id": original["revision_id"], "markdown": edited, "metadata": metadata}
    response = client.patch(f"/admin/versions/{vid}", json=body, headers=headers(client))
    assert response.status_code == 200
    assert client.patch(f"/admin/versions/{vid}", json=body, headers=headers(client)).status_code == 409
    revision = client.get(f"/admin/versions/{vid}/normalized").json()
    assert revision["markdown"] == edited
    assert revision["provenance"][-1]["origin"] == "manual"
    historical = client.get(f"/admin/versions/{vid}/normalized?revision_id={original['revision_id']}").json()
    assert historical["markdown"] == original["markdown"]
    assert historical["metadata"]["author"] == "" and revision["metadata"]["author"] == "Autor revisado"
    assert [r["number"] for r in client.get(f"/admin/versions/{vid}/revisions").json()["items"]] == [2, 1]
    assert client.patch(f"/admin/versions/{vid}", json=body | {"metadata": metadata | {"title": "   "}}, headers=headers(client)).status_code == 422
    assert client.post(f"/admin/versions/{vid}/review", json={"expected_revision_id": original["revision_id"]}, headers=headers(client)).status_code == 409
    approved = client.post(f"/admin/versions/{vid}/review", json={"expected_revision_id": revision["revision_id"]}, headers=headers(client))
    assert approved.status_code == 200 and approved.json()["reviewed_at"]
    assert approved.json()["status"] == "requiere_revision"


def test_empty_document_requires_review_and_cannot_be_approved(application):
    app, client, sessions = application
    login(client)
    data = upload(client, "empty.txt", b"", mime="text/plain").json()
    complete(app, sessions, data)
    version = client.get(f"/admin/documents/{data['document']['id']}").json()["versions"][0]
    assert version["status"] == "requiere_revision" and "empty_document" in version["diagnostics"]
    assert client.post(f"/admin/versions/{version['id']}/review", json={"expected_revision_id": version["revision_id"]}, headers=headers(client)).status_code == 422


def test_cancel_retry_lease_reclaim_and_no_late_writes(application):
    app, client, sessions = application
    login(client)
    data = upload(client).json()
    jid = data["job"]["id"]
    assert client.post(f"/admin/jobs/{jid}/cancel", headers=headers(client)).json()["status"] == "cancelado"
    assert claim(sessions, app.state.settings, "owner") is None
    retried = client.post(f"/admin/jobs/{jid}/retry", headers=headers(client)).json()
    assert retried["status"] == "pendiente" and retried["id"] != jid
    with sessions() as db:
        assert db.get(IngestionJob, jid).status == "cancelado"
    jid = retried["id"]
    old = claim(sessions, app.state.settings, "old-owner")
    with sessions() as db:
        db.get(IngestionJob, jid).lease_until = int(time.time()) - 1
        db.commit()
    new = claim(sessions, app.state.settings, "new-owner")
    assert old.generation < new.generation
    assert renew(sessions, app.state.settings, old) == "lost"
    assert not finalize_conversion(sessions, app.state.settings, old, {"markdown": "late", "provenance": [], "diagnostics": []})
    assert client.post(f"/admin/jobs/{jid}/cancel", headers=headers(client)).json()["status"] == "cancelando"
    assert renew(sessions, app.state.settings, new) == "cancel"
    finish_cancel(sessions, new)
    with sessions() as db:
        assert db.get(IngestionJob, jid).status == "cancelado"
        assert not db.scalars(select(NormalizedRevision)).all()


def test_replacement_preserves_versions_and_controlled_deletion(application):
    app, client, sessions = application
    login(client)
    data = upload(client).json()
    did = data["document"]["id"]
    assert client.request("DELETE", f"/admin/documents/{did}", json={"confirmation": data['document']['title']}, headers=headers(client)).status_code == 409
    complete(app, sessions, data)
    second = upload(client, "replacement.txt", b"Nuevo texto", path=f"/admin/documents/{did}/versions", mime="text/plain")
    assert second.status_code == 201
    assert [version["number"] for version in second.json()["document"]["versions"]] == [2, 1]
    complete(app, sessions, second.json())
    assert client.request("DELETE", f"/admin/documents/{did}", json={"confirmation": "incorrect"}, headers=headers(client)).status_code == 409
    response = client.request("DELETE", f"/admin/documents/{did}", json={"confirmation": data["document"]["title"]}, headers=headers(client))
    assert response.status_code == 202
    lease = claim(sessions, app.state.settings, "delete-owner")
    process_delete(sessions, app.state.settings, lease)
    assert client.get(f"/admin/documents/{did}").status_code == 404
    with sessions() as db:
        assert not db.scalars(select(DocumentFile)).all()
        assert not db.scalars(select(NormalizedRevision)).all()
        assert all(version.status == "eliminado" for version in db.scalars(select(DocumentVersion)))


def test_retention_and_withdrawal(application):
    app, client, sessions = application
    login(client)
    data = upload(client).json()
    complete(app, sessions, data)
    did, vid = data["document"]["id"], data["document"]["versions"][0]["id"]
    with sessions() as db:
        db.get(DocumentVersion, vid).status = "publicado"
        db.get(DocumentVersion, vid).published_at = int(time.time())
        db.get(Document, did).active_version_id = vid
        db.commit()
    assert client.post(f"/admin/documents/{did}/withdraw", headers=headers(client)).status_code == 200
    assert client.request("DELETE", f"/admin/documents/{did}", json={"confirmation": data["document"]["title"]}, headers=headers(client)).status_code == 409


@pytest.mark.parametrize("filename,content,mime", [
    ("../x.md", b"text", "text/markdown"), ("x.exe", b"text", "application/octet-stream"),
    ("x.pdf", b"not pdf", "application/pdf"), ("x.md", b"%PDF-1.7 fake", "text/markdown"),
    ("x.md", b"hello\x00", "text/markdown"), ("x.md", b"text", "application/pdf"),
    ("x.txt", b"\xff\xfe", "text/plain"), ("x.docx", b"zip", "application/zip")
])
def test_unsafe_uploads_leave_no_documents(application, filename, content, mime):
    _, client, sessions = application
    login(client)
    assert upload(client, filename, content, mime=mime).status_code == 422
    with sessions() as db:
        assert not db.scalars(select(Document)).all()


def test_size_limits_and_path_traversal(tmp_path):
    s = settings(tmp_path, upload_max_bytes=1024)
    with pytest.raises(InvalidFile, match="upload_too_large"):
        save_upload(io.BytesIO(b"x" * 1025), "x.txt", "text/plain", s)
    assert not list((tmp_path / "originals").iterdir())
    for key in ("../secret", "originals/../../secret", "C:\\secret", ""):
        with pytest.raises(InvalidFile):
            storage_path(s, key)


def test_docx_structure_and_pdf_page_provenance(tmp_path):
    import pymupdf
    s = settings(tmp_path)
    word = WordDocument()
    word.add_heading("Titulo", 1)
    word.add_paragraph("Texto en orden")
    table = word.add_table(rows=2, cols=2)
    for cell, value in zip([cell for row in table.rows for cell in row.cells], ["A", "B", "Uno", "Dos"]):
        cell.text = value
    word.add_paragraph("Primer paso", style="List Bullet")
    file = tmp_path / "sample.docx"
    word.save(file)
    result = convert(file, "docx", s)
    assert "# Titulo" in result["markdown"] and "| A | B |\n| --- | --- |" in result["markdown"]
    assert "- Primer paso" in result["markdown"]
    assert result["provenance"][-1]["section_path"] == ["Titulo"]
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "Pagina uno")
    pdf.new_page()
    file = tmp_path / "sample.pdf"
    pdf.save(file)
    pdf.close()
    result = convert(file, "pdf", s)
    assert "Pagina uno" in result["markdown"]
    assert result["provenance"][0]["page"] == 1
    assert "possible_ocr_page_2" in result["diagnostics"]


def test_manual_edits_do_not_keep_fabricated_pages():
    original = "# Titulo\n\nTexto original.\n"
    mapping = provenance(original, "pdf", page=4)
    changed = edited_provenance(original, mapping, "# Titulo\n\nTexto distinto.\n")
    assert changed[0]["page"] == 4
    assert changed[1]["page"] is None and changed[1]["origin"] == "manual"


@pytest.mark.parametrize("attack", ["traversal", "expansion", "entity", "corrupt_xml"])
def test_docx_archive_attacks_are_rejected(tmp_path, attack):
    buffer = io.BytesIO()
    types = '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
    document = '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/></w:document>'
    if attack == "entity":
        types = '<!DOCTYPE Types [<!ENTITY unsafe "test">]>' + types.replace("<Override", "&unsafe;<Override")
    if attack == "corrupt_xml":
        document = '<w:document'
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", types)
        archive.writestr("word/document.xml", document)
        if attack == "traversal":
            archive.writestr("../escape", "no")
        if attack == "expansion":
            archive.writestr("word/large.xml", "x" * 2000000)
    buffer.seek(0)
    with pytest.raises(InvalidFile):
        save_upload(buffer, "attack.docx", "application/zip", settings(tmp_path))
    assert not list((tmp_path / "originals").iterdir())


def test_worker_stop_preserves_original_and_retriable_job(application):
    app, client, sessions = application
    login(client)
    data = upload(client).json()
    lease = claim(sessions, app.state.settings, "stopping-worker")
    stop = threading.Event()
    stop.set()
    process(sessions, app.state.settings, lease, stop)
    with sessions() as db:
        job = db.get(IngestionJob, lease.id)
        assert job.status == "reintentable" and job.error_code == "worker_stopped"
        assert not db.scalars(select(NormalizedRevision)).all()
    assert client.get(f"/admin/versions/{data['document']['versions'][0]['id']}/original").content == b"# Titulo\n\nTexto original.\n"
