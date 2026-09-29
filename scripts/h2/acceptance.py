"""Aceptacion contra Compose con archivos sinteticos y sin imprimir credenciales."""
import io
import json
import subprocess
import time
import uuid
from pathlib import Path

import httpx
import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parents[2]


def compose(*args):
    result = subprocess.run(["docker", "compose", *args], cwd=ROOT, capture_output=True, text=True)
    if result.returncode:
        raise RuntimeError(f"Compose {args[0]} fallo")
    return result.stdout


def main():
    credentials = json.loads((ROOT / ".artifacts/h1/admin-credentials.json").read_text(encoding="utf-8"))
    report, created = {}, []
    origin = "http://localhost:3000"
    with httpx.Client(base_url=origin + "/api", timeout=60, trust_env=False) as client:
        client.post("/auth/login", json=credentials, headers={"Origin": origin}).raise_for_status()
        headers = {"Origin": origin, "X-CSRF-Token": client.get("/auth/me").json()["csrf_token"]}

        def request(method, path, **kwargs):
            response = client.request(method, path, headers=headers, **kwargs)
            response.raise_for_status()
            return response.json() if response.content else None

        def wait_job(jid):
            for _ in range(120):
                job = next(j for j in request("GET", "/admin/jobs")["items"] if j["id"] == jid)
                if job["status"] in {"completado", "fallido", "cancelado"}:
                    assert job["status"] == "completado", job
                    return job
                time.sleep(.5)
            raise RuntimeError("Trabajo no termino")

        def upload(name, content, mime, path="/admin/documents"):
            response = client.post(path, files={"file": (name, content, mime)},
                headers=headers | {"Idempotency-Key": str(uuid.uuid4())})
            response.raise_for_status()
            result = response.json()
            if result["document"]["id"] not in created:
                created.append(result["document"]["id"])
            return result

        pdf = pymupdf.open()
        pdf.new_page().insert_text((72, 72), "Texto PDF sintetico de aceptacion.")
        pdf_content = pdf.tobytes()
        pdf.close()
        word = Document()
        word.add_heading("Documento sintetico", 1)
        word.add_paragraph("Texto DOCX sintetico de aceptacion.")
        buffer = io.BytesIO()
        word.save(buffer)
        cases = [("sample.md", b"# Prueba\n\nTexto Markdown sintetico.\n", "text/markdown"),
                 ("sample.txt", b"Texto TXT sintetico.", "text/plain"),
                 ("sample.pdf", pdf_content, "application/pdf"),
                 ("sample.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")]
        try:
            for name, content, mime in cases:
                data = upload(name, content, mime)
                wait_job(data["job"]["id"])
                did, vid = data["document"]["id"], data["document"]["versions"][0]["id"]
                normalized = request("GET", f"/admin/versions/{vid}/normalized")
                assert normalized["markdown"].strip() and normalized["provenance"]
                assert client.get(f"/admin/versions/{vid}/original").content == content
                metadata = normalized["metadata"] | {"author": "Aceptacion sintetica"}
                request("PATCH", f"/admin/versions/{vid}", json={"expected_revision_id": normalized["revision_id"],
                    "markdown": normalized["markdown"] + "\nRevision manual.\n", "metadata": metadata})
                current = request("GET", f"/admin/versions/{vid}/normalized")
                assert current["provenance"][-1]["origin"] == "manual"
                assert request("GET", f"/admin/versions/{vid}/normalized?revision_id={normalized['revision_id']}")["metadata"]["author"] == ""
                reviewed = request("POST", f"/admin/versions/{vid}/review", json={"expected_revision_id": current["revision_id"]})
                assert reviewed["reviewed_at"] and reviewed["status"] == "requiere_revision"
                report[name] = {"original_intact": True, "revisions": 2, "reviewed": True,
                    "diagnostics": normalized["diagnostics"], "published": False}
            empty = upload("empty.txt", b"", "text/plain")
            wait_job(empty["job"]["id"])
            empty_version = request("GET", f"/admin/documents/{empty['document']['id']}")["versions"][0]
            assert "empty_document" in empty_version["diagnostics"]
            scanned = pymupdf.open()
            scanned.new_page()
            empty_pdf = upload("possible-scan.pdf", scanned.tobytes(), "application/pdf")
            scanned.close()
            wait_job(empty_pdf["job"]["id"])
            assert "possible_ocr_page_1" in request("GET", f"/admin/documents/{empty_pdf['document']['id']}")["versions"][0]["diagnostics"]
            invalid = client.post("/admin/documents", files={"file": ("broken.pdf", b"broken", "application/pdf")},
                headers=headers | {"Idempotency-Key": str(uuid.uuid4())})
            assert invalid.status_code == 422
            report["diagnostics"] = {"empty": True, "possible_scan": True, "corrupt_rejected": True}

            # Stop only this project's worker to exercise persisted cancellation and restart.
            compose("stop", "worker")
            try:
                queued = upload("restart.md", b"# Reinicio\n\nTrabajo persistente.", "text/markdown")
                jid = queued["job"]["id"]
                assert request("POST", f"/admin/jobs/{jid}/cancel")["status"] == "cancelado"
                retried = request("POST", f"/admin/jobs/{jid}/retry")
                assert retried["status"] == "pendiente" and retried["id"] != jid
                jid = retried["id"]
                # Claim as a crashed worker, then let its lease expire before restarting.
                code = """import time
from sqlalchemy import select
from bibliotecario.config import get_settings
from bibliotecario.db import build_database
from bibliotecario.jobs import claim
from bibliotecario.models import IngestionJob
s=get_settings(); engine,sessions=build_database(s)
lease=claim(sessions,s,'acceptance-crashed-worker')
assert lease is not None
with sessions() as db:
    db.get(IngestionJob,lease.id).lease_until=int(time.time())-1
    db.commit()
print(lease.generation)
"""
                assert compose("exec", "-T", "api", "python", "-c", code).strip() == "1"
            finally:
                compose("start", "worker")
            recovered = wait_job(jid)
            assert recovered["generation"] == 2 and recovered["attempts"] == 2
            events = request("GET", f"/admin/jobs/{jid}/events")["items"]
            assert any(event["event"] == "lease_expired_requeued" for event in events)
            assert [event["event"] for event in events] == ["queued", "manual_retry", "claimed", "lease_expired_requeued", "claimed", "converted"]
            stale_code = f"""from bibliotecario.config import get_settings
from bibliotecario.db import build_database
from bibliotecario.jobs import Lease,finalize_conversion
s=get_settings(); engine,sessions=build_database(s)
lease=Lease('{jid}',1,'acceptance-crashed-worker','convert')
assert not finalize_conversion(sessions,s,lease,{{'markdown':'Resultado tardio','provenance':[],'diagnostics':[]}})
print('stale_result_rejected')
"""
            assert compose("exec", "-T", "api", "python", "-c", stale_code).strip() == "stale_result_rejected"
            replacement = upload("replacement.txt", b"Nueva version sintetica.", "text/plain",
                f"/admin/documents/{queued['document']['id']}/versions")
            wait_job(replacement["job"]["id"])
            assert len(request("GET", f"/admin/documents/{queued['document']['id']}")["versions"]) == 2
            report["restart"] = {"cancel_retry": True, "expired_lease_recovered": True,
                "generation": recovered["generation"], "replacement_preserved_versions": True, "late_result_rejected": True}
        finally:
            for did in created:
                detail = client.get(f"/admin/documents/{did}")
                if detail.status_code == 200:
                    data = detail.json()
                    deletion = request("DELETE", f"/admin/documents/{did}", json={"confirmation": data["title"]})
                    wait_job(deletion["id"])
                    assert client.get(f"/admin/documents/{did}").status_code == 404
        report["synthetic_documents_deleted"] = len(created)
        request("POST", "/auth/logout")
    output = ROOT / ".artifacts/h2/runtime-acceptance.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
