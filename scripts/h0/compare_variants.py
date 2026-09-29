from __future__ import annotations

import argparse
import collections
import difflib
import sys
from pathlib import Path

from docx import Document
from docx.table import Table
from pypdf import PdfReader

from common import ROOT, clean_inline, load_json, normalize_text, sha256_file, word_tokens, write_json
from generate_variants import Block, blocks_from_markdown


def plain_source(blocks: list[Block]) -> str:
    parts: list[str] = []
    for block in blocks:
        if block.kind in {"heading", "paragraph", "list_item", "code"}:
            parts.append(block.text)
        elif block.kind == "table" and block.rows:
            parts.extend(" ".join(row) for row in block.rows)
    return "\n".join(parts)


def extract_docx(path: Path) -> tuple[str, dict]:
    document = Document(path)
    parts: list[str] = []
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            parts.extend(" ".join(cell.text for cell in row.cells) for row in item.rows)
        else:
            parts.append(item.text)
    return "\n".join(parts), {"paragraphs": len(document.paragraphs), "tables": len(document.tables)}


def extract_pdf(path: Path) -> tuple[str, dict]:
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages), {"pages": len(reader.pages)}


def multiset_recall(expected: list[str], actual: list[str]) -> float:
    if not expected:
        return 1.0
    expected_counts = collections.Counter(expected)
    actual_counts = collections.Counter(actual)
    matched = sum(min(count, actual_counts[token]) for token, count in expected_counts.items())
    return matched / sum(expected_counts.values())


def structural_anchors(blocks: list[Block]) -> list[str]:
    anchors: list[str] = []
    for block in blocks:
        if block.kind == "list_item" and block.text:
            anchors.append(block.text)
        elif block.kind == "code":
            anchors.extend(line.strip() for line in block.text.splitlines() if len(line.strip()) >= 6)
        elif block.kind == "table" and block.rows:
            anchors.extend(" ".join(cell for cell in row if cell) for row in block.rows)
    return anchors


def score(blocks: list[Block], extracted: str) -> dict:
    source_text = plain_source(blocks)
    source_tokens = word_tokens(source_text)
    extracted_tokens = word_tokens(extracted)
    normalized_extracted = normalize_text(extracted)
    extracted_token_text = " ".join(extracted_tokens)
    headings = [block.text for block in blocks if block.kind == "heading" and block.text]
    anchors = structural_anchors(blocks)
    return {
        "token_recall": round(multiset_recall(source_tokens, extracted_tokens), 4),
        "heading_recall": round(sum(" ".join(word_tokens(value)) in extracted_token_text for value in headings) / len(headings), 4) if headings else 1.0,
        "structure_recall": round(sum(normalize_text(value) in normalized_extracted for value in anchors) / len(anchors), 4) if anchors else 1.0,
        "order_ratio": round(difflib.SequenceMatcher(None, source_tokens, extracted_tokens, autojunk=False).ratio(), 4),
        "source_tokens": len(source_tokens),
        "extracted_tokens": len(extracted_tokens),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Compare generated variants with their Markdown sources.")
    parser.add_argument("--manifest", type=Path, default=ROOT / ".artifacts" / "h0" / "variants" / "manifest.json")
    parser.add_argument("--output", type=Path, default=ROOT / ".artifacts" / "h0" / "comparison.json")
    args = parser.parse_args()
    manifest = load_json(args.manifest)
    selection = load_json(ROOT / "evaluation" / "h0" / "variant_selection.json")
    report = {"schema_version": 1, "results": [], "passed": True}
    for variant in manifest["variants"]:
        source_path = ROOT / variant["source_path"]
        if sha256_file(source_path) != variant["source_sha256"]:
            raise ValueError(f"source changed for {variant['document_id']}")
        blocks = blocks_from_markdown(source_path.read_text(encoding="utf-8-sig"))
        for format_name, extractor in (("docx", extract_docx), ("pdf", extract_pdf)):
            output = ROOT / variant["outputs"][format_name]["path"]
            if sha256_file(output) != variant["outputs"][format_name]["sha256"]:
                raise ValueError(f"variant hash mismatch: {output}")
            extracted, inventory = extractor(output)
            metrics = score(blocks, extracted)
            thresholds = selection["thresholds"][format_name]
            gates = {name: metrics[name] >= minimum for name, minimum in thresholds.items()}
            passed = all(gates.values())
            report["passed"] = report["passed"] and passed
            report["results"].append({"document_id": variant["document_id"], "format": format_name, "passed": passed, "gates": gates, "metrics": metrics, "inventory": inventory, "extracted_sha256": __import__("hashlib").sha256(extracted.encode("utf-8")).hexdigest()})
            print(f"{variant['document_id']} {format_name}: {'PASS' if passed else 'FAIL'} {metrics}")
    write_json(args.output, report)
    print(f"comparison written to {args.output}")
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
