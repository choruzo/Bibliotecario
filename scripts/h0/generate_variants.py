from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import textwrap
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import LongTable, PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer, TableStyle
from xml.sax.saxutils import escape

from common import ROOT, catalog_by_id, clean_inline, load_json, sha256_file, write_json


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    ordered: bool = False
    rows: list[list[str]] | None = None


TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$")


def table_cells(line: str) -> list[str]:
    return [clean_inline(cell) for cell in line.strip().strip("|").split("|")]


def blocks_from_markdown(text: str) -> list[Block]:
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[Block] = []
    paragraph: list[str] = []
    index = 0

    def flush_paragraph() -> None:
        if paragraph:
            value = clean_inline(" ".join(part.strip() for part in paragraph))
            if value:
                blocks.append(Block("paragraph", value))
            paragraph.clear()

    while index < len(lines):
        line = lines[index]
        if line.lstrip().startswith("```"):
            flush_paragraph()
            fence = line.lstrip()[:3]
            language = line.lstrip()[3:].strip()
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].lstrip().startswith(fence):
                code.append(lines[index])
                index += 1
            blocks.append(Block("code", "\n".join(code), level=0, ordered=False))
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if heading:
            flush_paragraph()
            blocks.append(Block("heading", clean_inline(heading.group(2)), level=len(heading.group(1))))
            index += 1
            continue
        if "|" in line and index + 1 < len(lines) and TABLE_SEPARATOR.match(lines[index + 1]):
            flush_paragraph()
            rows = [table_cells(line)]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(table_cells(lines[index]))
                index += 1
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            blocks.append(Block("table", rows=rows))
            continue
        item = re.match(r"^\s*([-*+]|\d+[.)])\s+(.+)$", line)
        if item:
            flush_paragraph()
            blocks.append(Block("list_item", clean_inline(item.group(2)), ordered=item.group(1)[0].isdigit()))
            index += 1
            continue
        if not line.strip():
            flush_paragraph()
        elif re.match(r"^\s*(?:-{3,}|\*{3,}|_{3,})\s*$", line):
            flush_paragraph()
        else:
            paragraph.append(line)
        index += 1
    flush_paragraph()
    return blocks


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(cell, value: int = 100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def canonicalize_docx(path: Path) -> None:
    temporary = path.with_suffix(".canonical.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    os.replace(temporary, path)


def create_docx(blocks: list[Block], output: Path, document_id: str, source_hash: str) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    fixed = dt.datetime(2000, 1, 1, tzinfo=dt.timezone.utc)
    document.core_properties.title = document_id
    document.core_properties.subject = "H0 reproducible conversion fixture"
    document.core_properties.keywords = f"source-sha256:{source_hash}"
    document.core_properties.created = fixed
    document.core_properties.modified = fixed
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].font.size = Pt(10.5)
    first_heading = True
    for block in blocks:
        if block.kind == "heading":
            if first_heading and block.level == 1:
                paragraph = document.add_paragraph(block.text, style="Title")
                first_heading = False
            else:
                paragraph = document.add_heading(block.text, level=min(block.level, 4))
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(4)
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph(block.text)
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.line_spacing = 1.08
        elif block.kind == "list_item":
            style = "List Number" if block.ordered else "List Bullet"
            paragraph = document.add_paragraph(block.text, style=style)
            paragraph.paragraph_format.space_after = Pt(2)
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2)
            paragraph.paragraph_format.right_indent = Inches(0.15)
            paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(4)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F2F2F2")
            paragraph._p.get_or_add_pPr().append(shading)
            run = paragraph.add_run(block.text or " ")
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
            run.font.size = Pt(8.5)
        elif block.kind == "table" and block.rows:
            table = document.add_table(rows=len(block.rows), cols=len(block.rows[0]))
            table.style = "Table Grid"
            table.autofit = True
            set_repeat_table_header(table.rows[0])
            for row_index, row in enumerate(block.rows):
                for column_index, value in enumerate(row):
                    cell = table.cell(row_index, column_index)
                    cell.text = value
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_margins(cell)
                    if row_index == 0:
                        shade_cell(cell, "1F4E78")
                    elif row_index % 2 == 0:
                        shade_cell(cell, "EAF2F8")
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        for run in paragraph.runs:
                            run.font.name = "Arial"
                            run.font.size = Pt(8)
                            if row_index == 0:
                                run.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    canonicalize_docx(output)


def pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="H0Title", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=colors.black, spaceAfter=12, alignment=TA_LEFT))
    for level, size in ((1, 15), (2, 13), (3, 11), (4, 10)):
        styles.add(ParagraphStyle(name=f"H0H{level}", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=size, leading=size + 3, textColor=colors.black, spaceBefore=8, spaceAfter=4))
    styles.add(ParagraphStyle(name="H0Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=12, spaceAfter=5))
    styles.add(ParagraphStyle(name="H0List", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.2, leading=11.5, leftIndent=14, firstLineIndent=-8, spaceAfter=2))
    styles.add(ParagraphStyle(name="H0Cell", parent=styles["BodyText"], fontName="Helvetica", fontSize=6.8, leading=8.2))
    return styles


def create_pdf(blocks: list[Block], output: Path, document_id: str, source_hash: str) -> None:
    rl_config.invariant = 1
    styles = pdf_styles()
    story = []
    first_heading = True
    ordered_index = 0
    for block in blocks:
        if block.kind != "list_item":
            ordered_index = 0
        if block.kind == "heading":
            name = "H0Title" if first_heading and block.level == 1 else f"H0H{min(block.level, 4)}"
            first_heading = False
            story.append(Paragraph(escape(block.text), styles[name]))
        elif block.kind == "paragraph":
            story.append(Paragraph(escape(block.text), styles["H0Body"]))
        elif block.kind == "list_item":
            if block.ordered:
                ordered_index += 1
                marker = f"{ordered_index}."
            else:
                ordered_index = 0
                marker = "-"
            story.append(Paragraph(f"{marker} {escape(block.text)}", styles["H0List"]))
        elif block.kind == "code":
            wrapped = []
            for line in (block.text or " ").splitlines():
                wrapped.extend(textwrap.wrap(line, width=92, subsequent_indent="  ", replace_whitespace=False, drop_whitespace=False) or [""])
            story.append(Preformatted("\n".join(wrapped), ParagraphStyle(name="H0Code", fontName="Courier", fontSize=7.2, leading=9, leftIndent=8, rightIndent=8, backColor=colors.HexColor("#F2F2F2"), borderPadding=6, spaceBefore=4, spaceAfter=5)))
        elif block.kind == "table" and block.rows:
            column_count = len(block.rows[0])
            available = 7.0 * inch
            data = [[Paragraph(escape(value), styles["H0Cell"]) for value in row] for row in block.rows]
            table = LongTable(data, colWidths=[available / column_count] * column_count, repeatRows=1, hAlign="LEFT")
            commands = [
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9D9D9")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
            for row_index in range(2, len(block.rows), 2):
                commands.append(("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor("#EAF2F8")))
            table.setStyle(TableStyle(commands))
            story.extend([table, Spacer(1, 6)])
    output.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(str(output), pagesize=LETTER, rightMargin=0.75 * inch, leftMargin=0.75 * inch, topMargin=0.7 * inch, bottomMargin=0.7 * inch, title=document_id, subject=f"source-sha256:{source_hash}", author="Bibliotecario H0", invariant=1)
    document.build(story)


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate reproducible PDF and DOCX fixtures from selected Markdown sources.")
    parser.add_argument("--selection", type=Path, default=ROOT / "evaluation" / "h0" / "variant_selection.json")
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    selection = load_json(args.selection)
    output_dir = args.output_dir or ROOT / selection["output_directory"]
    catalog = catalog_by_id()
    manifest = {"schema_version": 1, "generator": "scripts/h0/generate_variants.py", "variants": []}
    for selected in selection["documents"]:
        entry = catalog[selected["document_id"]]
        source = ROOT / entry["path"]
        actual_hash = sha256_file(source)
        if actual_hash != entry["sha256"]:
            raise ValueError(f"source hash mismatch for {entry['id']}: {actual_hash}")
        blocks = blocks_from_markdown(source.read_text(encoding="utf-8-sig"))
        stem = entry["id"].lower()
        docx_path = output_dir / f"{stem}.docx"
        pdf_path = output_dir / f"{stem}.pdf"
        create_docx(blocks, docx_path, entry["id"], actual_hash)
        create_pdf(blocks, pdf_path, entry["id"], actual_hash)
        manifest["variants"].append({
            "document_id": entry["id"],
            "source_path": entry["path"],
            "source_sha256": actual_hash,
            "profile": selected["profile"],
            "outputs": {
                "docx": {"path": docx_path.relative_to(ROOT).as_posix(), "sha256": sha256_file(docx_path)},
                "pdf": {"path": pdf_path.relative_to(ROOT).as_posix(), "sha256": sha256_file(pdf_path)},
            },
        })
        print(f"generated {docx_path.name} and {pdf_path.name}")
    write_json(output_dir / "manifest.json", manifest)
    print(f"manifest written to {output_dir / 'manifest.json'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
